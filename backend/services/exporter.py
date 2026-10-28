# -*- coding: utf-8 -*-
"""导出服务：用例文档 Word/PDF/Excel，测试报告 HTML/PDF/Excel"""
import io
import os
import json
import base64
from datetime import datetime

import db as d
from utils import jloads

OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "exports")
os.makedirs(OUT_DIR, exist_ok=True)

# ---------- 中文字体支持（PDF） ----------
def _register_cjk():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    try:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        return 'STSong-Light'
    except Exception:
        return 'Helvetica'


def case_rows(pid):
    rows = d.q("""SELECT c.*, a.name AS api_name, a.method, a.path, a.module
                  FROM cases c LEFT JOIN apis a ON a.id=c.api_id WHERE c.project_id=?
                  ORDER BY a.module, c.api_id,
                  CASE c.case_type WHEN '正向' THEN 1 WHEN '反向' THEN 2 WHEN '边界' THEN 3 ELSE 4 END,
                  CASE c.priority WHEN 'P0' THEN 1 WHEN 'P1' THEN 2 ELSE 3 END, c.id""", (pid,))
    return [d.row_to_dict(r) for r in rows]


def _fmt_case(c):
    q = jloads(c["query_params"] if "query_params" in c.keys() else "{}", {})
    rd = jloads(c["request_data"], {})
    exp = jloads(c["expected"], {})
    asserts = []
    for a in (exp.get("assertions") or []):
        asserts.append(f"{a.get('type')}{(':' + a.get('expr')) if a.get('expr') else ''} {a.get('op')} {a.get('value')}")
    return {
        "id": c["id"], "name": c["name"], "api": c["api_name"] or "", "module": c["module"] or "",
        "method": c["method"] or "", "path": c["path"] or "",
        "type": c["case_type"], "priority": c["priority"], "steps": c["steps"] or "",
        "request": json.dumps(rd, ensure_ascii=False) if rd else "使用接口默认参数",
        "expected": f"状态码 {exp.get('status_code', '')}" + ("；" + "；".join(asserts) if asserts else ""),
        "created": c["created_at"],
    }


def export_cases_word(pid, project_name):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading(f"{project_name} - 接口测试用例文档", 0)
    doc.add_paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}　生成方式：AI 全自动生成 + 自动整理")
    tbl = doc.add_table(rows=1, cols=10)
    tbl.style = 'Light Grid Accent 1'
    hdr = tbl.rows[0].cells
    for i, h in enumerate(["用例ID", "用例名称", "所属接口", "模块", "类型", "优先级", "请求参数", "预期结果", "测试步骤", "创建时间"]):
        hdr[i].text = h
    for c in case_rows(pid):
        f = _fmt_case(c)
        cells = tbl.add_row().cells
        for i, v in enumerate([str(f["id"]), f["name"], f["api"], f["module"], f["type"], f["priority"],
                               f["request"], f["expected"], f["steps"], f["created"]]):
            cells[i].text = str(v)
    path = os.path.join(OUT_DIR, f"测试用例文档_{project_name}_{pid}.docx")
    doc.save(path)
    return path


def export_cases_excel(pid, project_name):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    headers = ["用例ID", "用例名称", "所属接口", "模块", "请求方式", "接口地址", "类型", "优先级", "请求参数", "预期结果", "测试步骤"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1677FF")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in case_rows(pid):
        f = _fmt_case(c)
        ws.append([f["id"], f["name"], f["api"], f["module"], f["method"], f["path"], f["type"],
                   f["priority"], f["request"], f["expected"], f["steps"]])
    for col in ws.columns:
        width = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width * 1.4 + 2, 10), 45)
    path = os.path.join(OUT_DIR, f"测试用例_{project_name}_{pid}.xlsx")
    wb.save(path)
    return path


def export_cases_pdf(pid, project_name):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    font = _register_cjk()
    style = ParagraphStyle('cjk', fontName=font, fontSize=8, leading=11)
    path = os.path.join(OUT_DIR, f"测试用例文档_{project_name}_{pid}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    story = [Paragraph(f"{project_name} - 接口测试用例文档", ParagraphStyle('t', fontName=font, fontSize=16, leading=22)),
             Spacer(1, 10)]
    data = [["用例ID", "用例名称", "所属接口", "类型", "优先级", "请求参数", "预期结果"]]
    for c in case_rows(pid):
        f = _fmt_case(c)
        data.append([str(f["id"]), Paragraph(f["name"], style), Paragraph(f["api"], style), f["type"],
                     f["priority"], Paragraph(f["request"], style), Paragraph(f["expected"], style)])
    tbl = Table(data, colWidths=[40, 110, 90, 40, 40, 130, 110], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1677FF")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E6EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
    ]))
    story.append(tbl)
    doc.build(story)
    return path


# ---------------- 测试报告导出 ----------------
def env_of_task(task: dict) -> dict:
    """任务执行环境快照（名称 + 域名），用于报告的环境可追溯性。
    环境已被删除时退化为仅保留名称，不阻断报告渲染。"""
    env = None
    if task.get("env_id"):
        env = d.q1("SELECT * FROM environments WHERE id=?", (task["env_id"],))
    if env is None and task.get("project_id") and task.get("env_name"):
        env = d.q1("SELECT * FROM environments WHERE project_id=? AND name=?",
                   (task["project_id"], task["env_name"]))
    env = d.row_to_dict(env) if env else None
    return {
        "env_name": (env or {}).get("name") or task.get("env_name") or "",
        "env_base_url": (env or {}).get("base_url") or "",
        "env_deleted": bool(task.get("env_name")) and env is None,
    }


def report_context(task_id):
    t = d.q1("SELECT * FROM tasks WHERE id=?", (task_id,))
    if not t:
        return None
    results = d.rows_to_dicts(d.q("SELECT * FROM case_results WHERE task_id=? ORDER BY id", (task_id,)))
    for r in results:
        for k in ("request_info", "response_info", "assertions"):
            r[k] = jloads(r[k], {} if k != "assertions" else [])
    task = d.row_to_dict(t)
    task.update(env_of_task(task))
    return {"task": task, "results": results}


def _pie_svg(passed, failed, skipped):
    total = max(passed + failed + skipped, 1)
    import math
    segs = [("#00B42A", passed), ("#F53F3F", failed), ("#86909C", skipped)]
    start, paths = -90.0, []
    for color, val in segs:
        if val <= 0:
            continue
        ang = val / total * 360
        end = start + ang
        x1, y1 = 60 + 45 * math.cos(math.radians(start)), 60 + 45 * math.sin(math.radians(start))
        x2, y2 = 60 + 45 * math.cos(math.radians(end)), 60 + 45 * math.sin(math.radians(end))
        large = 1 if ang > 180 else 0
        if ang >= 360:
            paths.append(f'<circle cx="60" cy="60" r="45" fill="{color}"/>')
        else:
            paths.append(f'<path d="M60 60 L{x1:.2f} {y1:.2f} A45 45 0 {large} 1 {x2:.2f} {y2:.2f} Z" fill="{color}"/>')
        start = end
    return "".join(paths)


def _bar_svg(module_stats):
    if not module_stats:
        return ""
    maxv = max(s[1] or 1 for s in module_stats)
    bars = ""
    for i, (mod, tot, okc) in enumerate(module_stats):
        h = 120 / max(len(module_stats), 1)
        y = i * (h + 8)
        rate = (okc / tot * 100) if tot else 0
        w = rate / 100 * 320
        bars += (f'<text x="0" y="{y + h / 2 + 4:.0f}" font-size="10" fill="#4E5969">{mod}</text>'
                 f'<rect x="90" y="{y:.0f}" width="320" height="{h:.0f}" fill="#F5F7FA" rx="3"/>'
                 f'<rect x="90" y="{y:.0f}" width="{w:.0f}" height="{h:.0f}" fill="#1677FF" rx="3"/>'
                 f'<text x="420" y="{y + h / 2 + 4:.0f}" font-size="10" fill="#4E5969">{rate:.0f}% ({okc}/{tot})</text>')
    return bars


def export_report_html(task_id):
    ctx = report_context(task_id)
    if not ctx:
        return None
    t, results = ctx["task"], ctx["results"]
    passed, failed, skipped = t["passed"], t["failed"], t["skipped"]
    total = max(t["total"], 1)
    rate = passed / max(t["executed"], 1) * 100 if t["executed"] else 0
    mod_map = {}
    for r in results:
        m = r["module"] or "默认模块"
        a, b, c = mod_map.get(m, (0, 0, 0))
        mod_map[m] = (a + 1, b + (1 if r["status"] == "passed" else 0), c)
    module_stats = [(m, v[0], v[1]) for m, v in mod_map.items()]
    rows = ""
    for r in results:
        color = {"passed": "#00B42A", "failed": "#F53F3F", "skipped": "#86909C"}.get(r["status"], "#86909C")
        label = {"passed": "通过", "failed": "失败", "skipped": "跳过"}.get(r["status"], r["status"])
        req = r.get("request_info") or {}
        url = req.get("url") or "-"
        rows += (f'<tr><td>{r["case_id"]}</td><td>{r["case_name"]}</td><td>{r["api_name"]}</td>'
                 f'<td>{r["case_type"]}</td><td class="url">{url}</td>'
                 f'<td style="color:{color};font-weight:600">{label}</td>'
                 f'<td>{r["duration"]}s</td><td class="reason">{r["fail_reason"] or "-"}</td></tr>')
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>{t['name']} 测试报告</title>
<style>
body{{font-family:"Microsoft YaHei","PingFang SC",sans-serif;background:#F5F7FA;color:#4E5969;margin:0;padding:24px}}
h1{{font-size:20px;color:#1D2129;margin:0 0 4px}}.sub{{font-size:12px;color:#86909C;margin-bottom:20px}}
.card{{background:#fff;border-radius:8px;padding:20px;box-shadow:0 2px 8px rgba(0,0,0,.06);margin-bottom:20px}}
.grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:20px}}
.stat{{background:#fff;border-radius:8px;padding:20px;box-shadow:0 2px 8px rgba(0,0,0,.06)}}
.stat .k{{font-size:14px;color:#86909C}}.stat .v{{font-size:26px;font-weight:700;color:#1D2129}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:20px}}
table{{width:100%;border-collapse:collapse;font-size:13px}}
th{{background:#F5F7FA;padding:10px;text-align:left;color:#1D2129;border-bottom:1px solid #E5E6EB}}
td{{padding:10px;border-bottom:1px solid #E5E6EB}}
.reason{{color:#F53F3F;font-size:12px}}
.url{{font-size:12px;color:#4E5969;word-break:break-all;max-width:280px}}
.env{{background:#fff;border-radius:8px;padding:14px 20px;box-shadow:0 2px 8px rgba(0,0,0,.06);margin-bottom:20px;font-size:13px}}
.env b{{color:#1D2129}}.env code{{background:#F5F7FA;padding:2px 8px;border-radius:4px;color:#1677FF;font-size:13px}}
h2{{font-size:16px;color:#1D2129;margin:0 0 12px}}
</style></head><body>
<h1>{t['name']} · 接口自动化测试报告</h1>
<div class="sub">执行时间：{t['created_at']}　任务耗时：{t['duration']}s　通过率：{rate:.1f}%</div>
<div class="env">执行环境：<b>{t.get('env_name') or '未指定'}</b>
{('　域名：<code>' + t['env_base_url'] + '</code>') if t.get('env_base_url') else '　<span style="color:#F53F3F">（未配置域名）</span>'}
{('<span style="color:#FF7D00">　注意：该环境已被删除，域名无法回溯</span>' if t.get('env_deleted') else '')}</div>
<div class="grid">
<div class="stat"><div class="k">总用例数</div><div class="v">{t['total']}</div></div>
<div class="stat"><div class="k">成功 / 失败 / 跳过</div><div class="v"><span style="color:#00B42A">{passed}</span> / <span style="color:#F53F3F">{failed}</span> / <span style="color:#86909C">{skipped}</span></div></div>
<div class="stat"><div class="k">通过率</div><div class="v">{rate:.1f}%</div></div>
</div>
<div class="two" style="margin-top:20px">
<div class="card"><h2>用例通过率分布</h2><svg width="120" height="120">{_pie_svg(passed, failed, skipped)}</svg>
<span style="font-size:12px;color:#00B42A">■ 通过 {passed}</span>
<span style="font-size:12px;color:#F53F3F">■ 失败 {failed}</span>
<span style="font-size:12px;color:#86909C">■ 跳过 {skipped}</span></div>
<div class="card"><h2>模块成功率</h2><svg width="500" height="{max(len(module_stats) * 28, 60)}">{_bar_svg(module_stats)}</svg></div>
</div>
<div class="card"><h2>用例执行明细</h2><table><thead><tr><th>用例ID</th><th>用例名称</th><th>所属接口</th><th>类型</th><th>请求地址</th><th>状态</th><th>耗时</th><th>失败原因</th></tr></thead><tbody>{rows or '<tr><td colspan="8">暂无数据</td></tr>'}</tbody></table></div>
</body></html>"""
    path = os.path.join(OUT_DIR, f"测试报告_{t['name']}.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path


def export_report_excel(task_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    ctx = report_context(task_id)
    if not ctx:
        return None
    t, results = ctx["task"], ctx["results"]
    wb = Workbook()
    ws = wb.active
    ws.title = "测试报告"
    ws.append(["任务名称", t["name"], "执行时间", t["created_at"]])
    ws.append(["执行环境", t.get("env_name") or "未指定", "环境域名", t.get("env_base_url") or "未配置"])
    ws.append(["总用例数", t["total"], "成功", t["passed"], "失败", t["failed"], "跳过", t["skipped"]])
    ws.append([])
    ws.append(["用例ID", "用例名称", "所属接口", "类型", "请求地址", "状态", "耗时(s)", "失败原因"])
    for cell in ws[5]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1677FF")
    for r in results:
        ws.append([r["case_id"], r["case_name"], r["api_name"], r["case_type"],
                   (r.get("request_info") or {}).get("url") or "-",
                   {"passed": "通过", "failed": "失败", "skipped": "跳过"}.get(r["status"]),
                   r["duration"], r["fail_reason"]])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width * 1.4 + 2, 10), 45)
    path = os.path.join(OUT_DIR, f"测试报告_{t['name']}.xlsx")
    wb.save(path)
    return path


def export_report_pdf(task_id):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    font = _register_cjk()
    ctx = report_context(task_id)
    if not ctx:
        return None
    t, results = ctx["task"], ctx["results"]
    style = ParagraphStyle('cjk', fontName=font, fontSize=8, leading=11)
    path = os.path.join(OUT_DIR, f"测试报告_{t['name']}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    rate = (t["passed"] / t["executed"] * 100) if t["executed"] else 0
    story = [Paragraph(f"{t['name']} · 接口自动化测试报告", ParagraphStyle('t', fontName=font, fontSize=16, leading=22)),
             Paragraph(f"执行时间：{t['created_at']}　耗时：{t['duration']}s　通过率：{rate:.1f}%",
                       ParagraphStyle('s', fontName=font, fontSize=10, textColor=colors.HexColor("#86909C"))),
             Paragraph(f"执行环境：<b>{t.get('env_name') or '未指定'}</b>　"
                       f"环境域名：<b>{t.get('env_base_url') or '未配置'}</b>",
                       ParagraphStyle('e', fontName=font, fontSize=10,
                                      textColor=colors.HexColor("#1677FF"))),
             Spacer(1, 12)]
    stat = [["总用例数", "成功", "失败", "跳过", "通过率", "耗时(s)"],
            [str(t["total"]), str(t["passed"]), str(t["failed"]), str(t["skipped"]), f"{rate:.1f}%", str(t["duration"])]]
    st = Table(stat, colWidths=[80, 80, 80, 80, 80, 80])
    st.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F3FF")),
                            ("FONTNAME", (0, 0), (-1, -1), font),
                            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E6EB")),
                            ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    story += [st, Spacer(1, 16)]
    data = [["用例ID", "用例名称", "所属接口", "类型", "请求地址", "状态", "耗时(s)", "失败原因"]]
    for r in results:
        label = {"passed": "通过", "failed": "失败", "skipped": "跳过"}.get(r["status"], "")
        data.append([str(r["case_id"]), Paragraph(r["case_name"], style), Paragraph(r["api_name"], style),
                     r["case_type"], Paragraph((r.get("request_info") or {}).get("url") or "-", style),
                     label, str(r["duration"]), Paragraph((r["fail_reason"] or "-")[:120], style)])
    tbl = Table(data, colWidths=[30, 85, 65, 32, 118, 30, 35, 85], repeatRows=1)
    tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1677FF")),
                             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                             ("FONTNAME", (0, 0), (-1, -1), font),
                             ("FONTSIZE", (0, 0), (-1, -1), 8),
                             ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E6EB")),
                             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")])]))
    story.append(tbl)
    doc.build(story)
    return path
